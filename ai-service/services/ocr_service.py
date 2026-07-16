import io
from pathlib import Path

def extract_text(file_content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(file_content)
        elif ext == ".docx":
            return _extract_docx(file_content)
        elif ext in (".xlsx", ".xls"):
            return _extract_excel(file_content)
        elif ext in (".png", ".jpg", ".jpeg", ".tiff"):
            return _extract_image(file_content)
        else:
            return file_content.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[Extraction partial - {str(e)[:100]}]"

def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def _extract_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def _extract_excel(content: bytes) -> str:
    # Try openpyxl, fallback to raw bytes decode
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        rows = []
        for sheet in wb.worksheets:
            rows.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(c) for c in row if c is not None)
                if row_text.strip():
                    rows.append(row_text)
        return "\n".join(rows)
    except Exception:
        return content.decode("utf-8", errors="ignore")[:2000]

def _extract_image(content: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(img)
    except Exception:
        return "[Image uploaded - OCR requires Tesseract installation]"
